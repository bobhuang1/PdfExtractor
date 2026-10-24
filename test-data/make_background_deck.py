#!/usr/bin/env python3
"""Author `background_deck.pptx`: a deck whose first slide has a picture
background.

A slide's background is not a shape in its shape tree. PowerPoint states it in
`p:cSld/p:bg/p:bgPr`, and a photographic one is an `a:blipFill` there, naming a
media part through the slide's own relationships:

    <p:bg>
      <p:bgPr>
        <a:blipFill rotWithShape="1">
          <a:blip r:embed="rIdBackdrop"/>
          <a:stretch><a:fillRect/></a:stretch>
        </a:blipFill>
        <a:effectLst/>
      </p:bgPr>
    </p:bg>

python-pptx can author a slide's shapes but not a picture fill on its
background, so the deck is written with two slides and the first slide's
background is completed in the saved package: the photograph (`images/scan.jpg`,
the 620 x 350 scan the other office fixtures use), a relationship naming it, and
the `p:bg` element before the shape tree. The second slide has no `p:bg` at all,
so the test can hold the reader to a per-slide fact rather than a deck-wide one.

The first slide also carries a sentence, so the words in front of the backdrop
are read beside it. The second slide says its own sentence.

Regenerate with:

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_background_deck.py

The deck is written by python-pptx — a real PPTX writer — and the background is
the one piece it cannot author, added the way PowerPoint writes it. The tests
pin what the reader makes of the package, never the container's timestamps.
"""

from pathlib import Path
import zipfile

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
OUT = HERE / "background_deck.pptx"
BACKDROP = HERE / "images" / "scan.jpg"

SLIDE = "ppt/slides/slide1.xml"
RELS = "ppt/slides/_rels/slide1.xml.rels"
MEDIA = "ppt/media/backdrop.jpeg"
REL_ID = "rIdBackdrop"
REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


def main() -> None:
    deck = Presentation()

    first = deck.slides.add_slide(deck.slide_layouts[6])  # blank
    box = first.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(6.0), Inches(0.6))
    box.text_frame.text = "A photograph behind the words."

    second = deck.slides.add_slide(deck.slide_layouts[6])
    box = second.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(6.0), Inches(0.6))
    box.text_frame.text = "No picture behind these words."

    deck.save(OUT)

    with zipfile.ZipFile(OUT) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    if MEDIA in parts:
        raise SystemExit(f"{MEDIA} is already in the deck")
    rels = parts[RELS].decode("utf-8")
    if REL_ID in rels:
        raise SystemExit(f"{RELS} already names {REL_ID}")
    parts[RELS] = rels.replace(
        "</Relationships>",
        f'<Relationship Id="{REL_ID}" Type="{REL_TYPE}" Target="../media/backdrop.jpeg"/>'
        "</Relationships>",
    ).encode("utf-8")
    parts[MEDIA] = BACKDROP.read_bytes()

    # The package already declares `jpeg` (python-pptx's default), but keep the
    # fixture honest if that changes: the part is a JPEG and must say so.
    types = parts["[Content_Types].xml"].decode("utf-8")
    if 'Extension="jpeg"' not in types:
        types = types.replace(
            "</Types>",
            '<Default Extension="jpeg" ContentType="image/jpeg"/></Types>',
        )
        parts["[Content_Types].xml"] = types.encode("utf-8")

    slide = parts[SLIDE].decode("utf-8")
    background = (
        "<p:bg><p:bgPr>"
        f'<a:blipFill rotWithShape="1"><a:blip r:embed="{REL_ID}"/>'
        "<a:stretch><a:fillRect/></a:stretch></a:blipFill>"
        "<a:effectLst/>"
        "</p:bgPr></p:bg>"
    )
    if "<p:bg>" in slide:
        raise SystemExit(f"{SLIDE} already has a background")
    marker = "<p:cSld>"
    if marker not in slide:
        raise SystemExit(f"{SLIDE} has no p:cSld")
    parts[SLIDE] = slide.replace(marker, marker + background, 1).encode("utf-8")

    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)

    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
