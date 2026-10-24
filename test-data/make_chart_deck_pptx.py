#!/usr/bin/env python3
"""Author `chart_deck.pptx`, the deck whose only picture is a drawing.

A chart is a picture with no picture file behind it: the slide's `graphicFrame`
names a chart part, whose caches hold the title, the categories and the series,
and there is no media part anywhere in the package. The document model says so
with a `PictureItem` that has a `tabular_chart` and no `image` at all, which is
the one shape that makes the renderers' picture handling interesting:

* `referenced` writes nothing for it — there is no file to point at, and a
  broken image link is worse than none;
* `embedded` has nothing to carry, so it marks the spot instead of claiming
  bytes;
* `placeholder` marks the spot, as it does for every picture.

That is the document the placeholder gate needs: every golden in the set has
pictures with files, so a golden whose only picture has none is the one that
reaches the branch where a mark has no referenced block to stand in for.

The deck is deliberately small — a title, one paragraph, one chart — so the
goldens beside it can be read in full and the difference between the referenced
page and the placeholder page is one mark and nothing else.

    ~/.venvs/pdfextractor-docling/bin/python make_chart_deck_pptx.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "chart_deck.pptx"


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "A drawing with no picture file"

    # A text box, not the content placeholder: the words beside the drawing
    # should be one paragraph, so the page around the mark reads as prose.
    words = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(1.0))
    words.text_frame.text = "The chart below is a drawing of its data, not an image part."
    for run in words.text_frame.paragraphs[0].runs:
        run.font.size = Pt(20)

    # Two categories, two series, one point missing: a gap the chart draws and
    # the reconstructed table keeps as a gap rather than a plotted zero.
    data = CategoryChartData()
    data.categories = ["Measured", "Predicted"]
    data.add_series("Before", (12, 18))
    data.add_series("After", (15, None))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(2.9),
        Inches(8.4),
        Inches(3.8),
        data,
    )
    chart.chart.has_title = True
    chart.chart.chart_title.text_frame.text = "Units by run"

    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
