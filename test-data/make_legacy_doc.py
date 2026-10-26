#!/usr/bin/env python3
"""Write `test-data/brief.docx` and `test-data/legacy_brief.doc`.

The pair exists for the Writer side of the LibreOffice bridge, the way
`legacy_deck.ppt` and `q3_review.pptx` exist for the presentation side and
`ledger.xls` / `ledger.xlsx` for the spreadsheet one: the legacy file is the
native one carried through a real legacy writer, so a test can hold the whole
bridged document to the native reading and fail on anything the conversion
changed. The native document is written here with `python-docx`, and the `.doc`
is that document converted by LibreOffice's own "MS Word 97" filter — the same
importer and exporter the bridge itself uses.

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_legacy_doc.py

Needs LibreOffice (`soffice` on `PATH`, or `PDFE_SOFFICE` naming it), the same
dependency the bridge itself has.

What the pair carries, one case per feature the DOCX reader has:

* `Quarterly Brief` — a heading, read as a `section_header`.
* A paragraph whose words are split by a **hyperlink** (`The filing`), which the
  reader keeps as its own item with the URL in `hyperlink`.
* A bulleted pair and a numbered pair, as `list_item`s with their own markers.
  The numbering is put on each paragraph (`w:numPr`) rather than left to the
  style, because that is how Word writes a list and what the reader resolves —
  a style-only list would read without markers on the native side and with them
  on the bridged one, a difference of the fixture rather than of the bridge.
* A 3 x 3 table with a **column merge** (`Revenue` spanning two columns) and a
  **row merge** (`North`/`South`), so the bridge has to carry `gridSpan`,
  `vMerge` and the grid's width, not only a rectangle of words.
* A **picture** — `test-data/images/scan.jpg` — so the media part, the media
  type, the box and the pixel size all have to cross. Word 97 stores the image
  in the binary document and LibreOffice writes it back into the package it
  produces, which is exactly the sort of part a bridge can lose.

Everything here is a feature the native reading and the bridged one agree on
exactly, so the gate needs no excuse but the container's own identity and the
directory an output names its pictures in.

The script also writes `test-data/stamp_fields.doc`, which has no native twin
because its native and bridged readings are meant to *differ*: its footer holds a
`DATE` field and a `TIME` field, and every conversion rewrites both from its own
clock. LibreOffice reads the fields out of the Word 97 file and writes the moment
*it* converted into the package, so a reading of that `.doc` holds today's date and
now's time while any other conversion of the same file holds its own — which is
what `without_calendar_values` (either copy) keeps out of the character share, so
the ratchet cannot read the calendar or the clock. The stored results (`1/1/2000`,
`00:00:00`) never survive a conversion; they are there so the fields are valid
before one happens.
"""

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

HERE = Path(__file__).resolve().parent
NATIVE = HERE / "brief.docx"
BRIDGED = HERE / "legacy_brief.doc"
STAMPED = HERE / "stamp_fields.doc"
PICTURE = HERE / "images" / "scan.jpg"


def soffice() -> str:
    explicit = os.environ.get("PDFE_SOFFICE")
    if explicit:
        return explicit
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if not found:
        sys.exit("no soffice on PATH and PDFE_SOFFICE is not set")
    return found


def list_number(document, style_name: str) -> str:
    """The `w:numId` a list style already names, for the paragraph to state."""
    style = document.styles[style_name].element
    num_pr = style.find(qn("w:pPr")).find(qn("w:numPr"))
    return num_pr.find(qn("w:numId")).get(qn("w:val"))


def add_list(document, text: str, style_name: str) -> None:
    """A list item whose numbering sits on the paragraph, as Word writes one."""
    paragraph = document.add_paragraph(text, style=style_name)
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num_pr.append(level)
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), list_number(document, style_name))
    num_pr.append(number)
    paragraph._p.get_or_add_pPr().append(num_pr)


def add_link(paragraph, text: str, url: str) -> None:
    """A `w:hyperlink` around its own run, pointing at an external target."""
    relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    words = OxmlElement("w:t")
    words.text = text
    run.append(words)
    link.append(run)
    paragraph._p.append(link)


def literal(paragraph, text: str) -> None:
    """A run of words as written, whitespace and all."""
    run = OxmlElement("w:r")
    words = OxmlElement("w:t")
    words.set(qn("xml:space"), "preserve")
    words.text = text
    run.append(words)
    paragraph._p.append(run)


def field_mark(paragraph, kind: str) -> None:
    """One of a field's three marks: `begin`, `separate` or `end`."""
    run = OxmlElement("w:r")
    mark = OxmlElement("w:fldChar")
    mark.set(qn("w:fldCharType"), kind)
    run.append(mark)
    paragraph._p.append(run)


def add_field(paragraph, instruction: str, stored: str) -> None:
    """A field: `begin`, its instruction, `separate`, the stored result, `end`.

    The stored result is what a reader that does not refresh fields holds; the
    conversion that *does* refresh them writes its own over it.
    """
    field_mark(paragraph, "begin")
    run = OxmlElement("w:r")
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = f" {instruction} "
    run.append(instruction_text)
    paragraph._p.append(run)
    field_mark(paragraph, "separate")
    literal(paragraph, stored)
    field_mark(paragraph, "end")


def write_stamp_fixture(scratch: Path) -> None:
    """The `.doc` whose footer is stamped by whichever conversion writes it."""
    document = Document()
    document.add_paragraph(
        "The footer says when this file was printed. The fields in it are refreshed "
        "by whatever writes the file next, so no reading of it holds what it stores."
    )
    footer = document.sections[0].footer.paragraphs[0]
    literal(footer, "Printed ")
    add_field(footer, 'DATE \\@ "M/d/yyyy"', "1/1/2000")
    literal(footer, " at ")
    add_field(footer, 'TIME \\@ "HH:mm:ss"', "00:00:00")
    native = scratch / "stamp_fields.docx"
    document.save(native)

    shutil.copyfile(convert(native, scratch), STAMPED)
    print(f"wrote {STAMPED}")


def write_native() -> None:
    document = Document()
    document.add_heading("Quarterly Brief", level=0)
    document.add_heading("Summary", level=1)
    opening = document.add_paragraph("The quarter closed ahead of plan. ")
    add_link(opening, "The filing", "https://example.com/filing")
    opening.add_run(" explains why.")
    add_list(document, "First bullet", "List Bullet")
    add_list(document, "Second bullet", "List Bullet")
    add_list(document, "Numbered one", "List Number")
    add_list(document, "Numbered two", "List Number")

    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(0, 1).merge(table.cell(0, 2))
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "42"
    table.cell(1, 2).text = "7"
    table.cell(2, 0).text = "South"
    table.cell(2, 1).text = "18"
    table.cell(2, 2).text = "3"
    table.cell(1, 0).merge(table.cell(2, 0))

    document.add_heading("Picture", level=2)
    document.add_paragraph("The chart below is the whole story.")
    document.add_picture(str(PICTURE), width=Inches(2))
    document.add_paragraph("Supply is tight.")
    document.save(NATIVE)


def convert(native: Path, scratch: Path) -> Path:
    """The Word 97 file LibreOffice's exporter writes for `native`."""
    out = scratch / "out"
    out.mkdir(exist_ok=True)
    profile = scratch / "profile"
    result = subprocess.run(
        [
            soffice(),
            "--headless",
            "--norestore",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to",
            "doc",
            "--outdir",
            str(out),
            str(native),
        ],
        capture_output=True,
        text=True,
    )
    produced = out / f"{native.stem}.doc"
    if result.returncode != 0 or not produced.is_file():
        sys.exit(
            "LibreOffice could not write the fixture: "
            f"{result.stderr.strip() or result.stdout.strip() or 'no diagnostic'}"
        )
    return produced


def main() -> None:
    write_native()

    with tempfile.TemporaryDirectory(prefix="pdfe-legacy-doc-") as scratch:
        produced = convert(NATIVE, Path(scratch))
        shutil.copyfile(produced, BRIDGED)
    print(f"wrote {NATIVE} and {BRIDGED}")

    with tempfile.TemporaryDirectory(prefix="pdfe-stamp-fields-") as scratch:
        write_stamp_fixture(Path(scratch))


if __name__ == "__main__":
    main()
