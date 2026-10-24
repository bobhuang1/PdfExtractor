#!/usr/bin/env python3
"""Write `gif_svg_pictures.docx`, `gif_svg_pictures.pptx` and their two pictures.

A picture is not always a PNG or a JPEG. Word and PowerPoint store a GIF as a
GIF — the grid is in its own header — and a vector drawing as an SVG whose root
tag states its size. The DOCX and PPTX readers take a media part's type from the
shared `picture-formats` list, which carries both, so either encoding must leave
a conversion as a picture: its media type, its pixel size, and its bytes copied
beside the JSON. These fixtures put both encodings in front of that rule, in a
Word and in a PowerPoint package, so the test can fail if a package stops being
read as pictures at all, or if only the raster one survives.

Each package carries the same two pictures in the same order, with a sentence
before and after them:

* `images/chart.gif` — a 60 × 36 bar chart, written by Pillow in the GIF format
  its own header describes;
* `images/diagram.svg` — a 40 × 24 line drawing, whose `width`/`height` on the
  root tag are the only place its size exists.

Neither office library can place an SVG: python-docx and python-pptx both take
a picture from a raster file and reject an SVG. The SVG is therefore written
into the saved package the way LibreOffice writes one — a media part named
`image2.svg`, the picture's own relationship pointing at it, and a
content-type declaration for the extension — after the library has written the
same picture as a one-pixel PNG stub. (Word 2016 and later write the picture
the other way round, a PNG fallback with the drawing in an `asvg:svgBlip`
beside it; `make_svg_blip_pictures.py` builds that pair.) The GIF travels
through the library untouched, which is what makes it a real GIF in the package
rather than a PNG with a new name.

Regenerate with:

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_gif_svg_pictures.py

The libraries are used only to write the fixtures; the readers under test never
see this script, and they read the committed bytes.

An office application is asked whether the assembled package is one it accepts:

    cargo test -p pdfextractor-cli --test cli_test \
        libreoffice_opens_and_resaves_the_svg_fixture_packages

opens both files with headless LibreOffice and writes them back, where LibreOffice
is installed, and holds what it wrote to the drawing — the SVG part kept byte for
byte, still declared and still reached by a relationship, and a raster drawn from
its `<rect>` and `<polyline>` beside it as the fallback Word's SVG scheme carries.
A part LibreOffice cannot read leaves an empty raster or is dropped outright, so
the check fails rather than passing on the file's presence. That test is the one
thing here that needs an office application; the fixtures themselves do not.
"""

import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as PptxInches

HERE = Path(__file__).resolve().parent
IMAGES = HERE / "images"
GIF = IMAGES / "chart.gif"
SVG = IMAGES / "diagram.svg"
DOCX = HERE / "gif_svg_pictures.docx"
PPTX = HERE / "gif_svg_pictures.pptx"

GIF_SIZE = (60, 36)
SVG_SIZE = (40, 24)

# The SVG's root tag is where a reader without an image decoder finds its size.
SVG_BYTES = (
    '<svg xmlns="http://www.w3.org/2000/svg" width="40" height="24" '
    'viewBox="0 0 40 24">'
    '<rect width="40" height="24" rx="3" fill="#e8eef7"/>'
    '<polyline points="4,19 14,11 22,15 36,5" fill="none" '
    'stroke="#2f6db5" stroke-width="2"/>'
    "</svg>"
).encode("utf-8")

# A one-pixel PNG, written only so the office library has something to place:
# its part is replaced by the SVG before the fixture is saved.
STUB = HERE / ".svg-stub.png"


def write_pictures() -> None:
    IMAGES.mkdir(exist_ok=True)

    chart = Image.new("RGB", GIF_SIZE, (232, 238, 247))
    draw = ImageDraw.Draw(chart)
    for top, colour in [(6, "#2f6db5"), (16, "#6da33a"), (26, "#e0b23b")]:
        draw.rectangle([(4, top), (55, top + 6)], fill=colour)
    chart.save(GIF)

    SVG.write_bytes(SVG_BYTES)

    Image.new("RGB", (1, 1), "white").save(STUB)


def write_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Before the pictures.")
    document.add_picture(str(GIF), width=Inches(1.5), height=Inches(0.9))
    document.add_picture(str(STUB), width=Inches(1.0), height=Inches(0.6))
    document.add_paragraph("After the pictures.")
    document.save(path)


def write_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank

    before = slide.shapes.add_textbox(
        PptxInches(0.8), PptxInches(0.4), PptxInches(4.0), PptxInches(0.4)
    )
    before.text_frame.text = "Before the pictures."

    slide.shapes.add_picture(
        str(GIF),
        PptxInches(0.8),
        PptxInches(1.0),
        width=PptxInches(1.5),
        height=PptxInches(0.9),
    )
    slide.shapes.add_picture(
        str(STUB),
        PptxInches(0.8),
        PptxInches(2.2),
        width=PptxInches(1.0),
        height=PptxInches(0.6),
    )

    after = slide.shapes.add_textbox(
        PptxInches(0.8), PptxInches(3.2), PptxInches(4.0), PptxInches(0.4)
    )
    after.text_frame.text = "After the pictures."

    presentation.save(path)


def replace_stub_with_svg(path: Path) -> None:
    """Rename the stub's media part and point every relationship at the SVG.

    The office libraries name the parts by order (`image1.gif` for the picture
    added first, `image2.png` for the stub). That order is asserted rather than
    guessed: if a future library names them differently, the fixture is rebuilt
    deliberately instead of silently keeping a PNG where the test says SVG.
    """
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    stub = "image2.png"
    svg = "image2.svg"
    media = [name for name in parts if name.endswith("/" + stub)]
    assert len(media) == 1, f"{path}: expected one {stub}, found {media}"
    media = media[0]
    svg_part = media[: -len(stub)] + svg

    parts[svg_part] = SVG_BYTES
    del parts[media]

    # Only relationships name a media part; the drawing itself refers to the
    # relationship, not to the file.
    for name in [name for name in parts if name.endswith(".rels")]:
        text = parts[name].decode("utf-8")
        parts[name] = text.replace(stub, svg).encode("utf-8")

    types = parts["[Content_Types].xml"].decode("utf-8")
    if 'Extension="svg"' not in types:
        marker = '<Default Extension="rels"'
        assert marker in types, f"{path}: no rels default to insert beside"
        types = types.replace(
            marker,
            '<Default Extension="svg" ContentType="image/svg+xml"/>' + marker,
        )
    parts["[Content_Types].xml"] = types.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def main() -> None:
    write_pictures()
    write_docx(DOCX)
    write_pptx(PPTX)
    replace_stub_with_svg(DOCX)
    replace_stub_with_svg(PPTX)
    STUB.unlink()
    for path in (GIF, SVG, DOCX, PPTX):
        print(f"wrote {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
