#!/usr/bin/env python3
"""Write `svg_blip_pictures.docx` and `svg_blip_pictures.pptx`.

Word and PowerPoint 2016 and later store an SVG picture twice. The picture's
`a:blip` keeps a PNG fallback in its own `r:embed`, for readers that cannot draw
SVG, and the drawing itself is named in an `asvg:svgBlip` extension beside it:

    <a:blip r:embed="rId5">
      <a:extLst><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">
        <asvg:svgBlip r:embed="rId6"/>
      </a:ext></a:extLst>
    </a:blip>

The fallback is only what an old reader shows; the SVG is the picture the file
asks for. A reader that stops at `r:embed` reports a raster of a vector drawing
— a picture of the wrong thing, at a resolution the file never chose.

Neither python-docx nor python-pptx can place an SVG, so each package is written
with a PNG in the picture's place and then completed the way Word completes it:
a media part for the drawing, a relationship naming it, a content-type
declaration for the extension, and the `asvg:svgBlip` extension in the drawing
part. The fallback PNG is drawn as a grey box, deliberately unlike the drawing,
so a reader that takes it is visible in what comes out.

The drawing is `images/diagram.svg`, the same 40 × 24 one the GIF/SVG fixtures
use.

The same drawing has another home in a package. LibreOffice writes an SVG where
it writes any picture — the picture's own `r:embed` names it, with no raster
beside it — and that is the spelling `test-data/gif_svg_pictures.*` carries. It
is the same picture at the same place, so the two spellings have to convert to
the same document; `svg_embed_pictures.docx` and `svg_embed_pictures.pptx` are
written here, from the same sentences and the same drawing, for one end-to-end
test to hold the two spellings' JSON and HTML to equality (the input's hash
excepted, which is a property of the bytes and the one thing that must differ).

Regenerate with:

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_svg_blip_pictures.py

The libraries are used only to write the fixtures; the readers under test never
see this script, and they read the committed bytes.
"""

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as PptxInches

HERE = Path(__file__).resolve().parent
IMAGES = HERE / "images"
SVG = IMAGES / "diagram.svg"
DOCX = HERE / "svg_blip_pictures.docx"
PPTX = HERE / "svg_blip_pictures.pptx"
# The same document and the same drawing, stored the way LibreOffice stores a
# picture: the SVG is the picture's own media part, with no raster beside it.
EMBED_DOCX = HERE / "svg_embed_pictures.docx"
EMBED_PPTX = HERE / "svg_embed_pictures.pptx"

# A grey box: what a raster fallback of the drawing would look like if nobody
# had drawn the drawing. Nothing about it is the SVG's pale fill or blue line.
FALLBACK_SIZE = (40, 24)
FALLBACK = HERE / ".svg-blip-fallback.png"

SVG_EXTENSION_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"


def write_fallback() -> None:
    box = Image.new("RGB", FALLBACK_SIZE, (128, 128, 128))
    draw = ImageDraw.Draw(box)
    draw.rectangle([(2, 2), (37, 21)], outline=(64, 64, 64))
    box.save(FALLBACK)


def write_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Before the picture.")
    document.add_picture(str(FALLBACK), width=Inches(1.0), height=Inches(0.6))
    document.add_paragraph("After the picture.")
    document.save(path)


def write_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank

    before = slide.shapes.add_textbox(
        PptxInches(0.8), PptxInches(0.4), PptxInches(4.0), PptxInches(0.4)
    )
    before.text_frame.text = "Before the picture."

    slide.shapes.add_picture(
        str(FALLBACK),
        PptxInches(0.8),
        PptxInches(1.0),
        width=PptxInches(1.0),
        height=PptxInches(0.6),
    )

    after = slide.shapes.add_textbox(
        PptxInches(0.8), PptxInches(2.2), PptxInches(4.0), PptxInches(0.4)
    )
    after.text_frame.text = "After the picture."

    presentation.save(path)


def source(svg_part: str) -> str:
    """Where the drawing sits, as the drawing part's own relationships say it."""
    assert svg_part.startswith(("word/", "ppt/")), svg_part
    if svg_part.startswith("word/"):
        return svg_part[len("word/") :]
    # A slide part's relationships resolve against `ppt/slides/`.
    return "../" + svg_part[len("ppt/") :]


def complete_with_svg(
    path: Path,
    drawing_part: str,
    rels_part: str,
    fallback_part: str,
    svg_part: str,
) -> None:
    """Add the drawing beside the fallback, the way Word states an SVG picture."""
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    rels = parts[rels_part].decode("utf-8")
    fallback = fallback_part.rsplit("/", 1)[1]
    # The one relationship that names the fallback is the picture's own.
    named = [
        (match.group(1), match.group(2))
        for match in re.finditer(
            r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels
        )
    ]
    matches = [entry for entry in named if entry[1].endswith("/" + fallback)]
    assert len(matches) == 1, f"{path}: expected one relationship to {fallback}, {named}"
    fallback_rel = matches[0][0]

    used = [int(match.group(1)) for match in re.finditer(r'Id="rId(\d+)"', rels)]
    svg_rel = f"rId{max(used) + 1}"
    parts[rels_part] = rels.replace(
        "</Relationships>",
        f'<Relationship Id="{svg_rel}" '
        f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
        f'Target="{source(svg_part)}"/></Relationships>',
    ).encode("utf-8")

    parts[svg_part] = SVG.read_bytes()

    drawing = parts[drawing_part].decode("utf-8")
    empty = f'<a:blip r:embed="{fallback_rel}"/>'
    assert drawing.count(empty) == 1, f"{path}: {empty} is not the picture's blip"
    parts[drawing_part] = drawing.replace(
        empty,
        f'<a:blip r:embed="{fallback_rel}">'
        f'<a:extLst><a:ext uri="{SVG_EXTENSION_URI}">'
        f'<asvg:svgBlip xmlns:asvg="{ASVG}" r:embed="{svg_rel}"/>'
        f"</a:ext></a:extLst></a:blip>",
    ).encode("utf-8")

    declare_svg_type(parts, path)
    rewrite(path, parts)


def embed_svg(path: Path, rels_part: str, fallback_part: str, svg_part: str) -> None:
    """Name the drawing where the picture itself points, the way LibreOffice
    stores an SVG picture: no raster fallback, and no extension beside it."""
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    rels = parts[rels_part].decode("utf-8")
    fallback = fallback_part.rsplit("/", 1)[1]
    named = [
        (match.group(1), match.group(2))
        for match in re.finditer(
            r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels
        )
    ]
    matches = [entry for entry in named if entry[1].endswith("/" + fallback)]
    assert len(matches) == 1, f"{path}: expected one relationship to {fallback}, {named}"
    picture_rel, _ = matches[0]
    parts[rels_part] = re.sub(
        rf'(<Relationship Id="{picture_rel}"[^>]*Target=")[^"]*(")',
        lambda match: match.group(1) + source(svg_part) + match.group(2),
        rels,
        count=1,
    ).encode("utf-8")

    parts[svg_part] = SVG.read_bytes()
    # The raster the picture was written with is no longer named by anything.
    del parts[fallback_part]
    declare_svg_type(parts, path)
    rewrite(path, parts)


def declare_svg_type(parts: dict[str, bytes], path: Path) -> None:
    """Declare the `svg` extension's media type, once."""
    types = parts["[Content_Types].xml"].decode("utf-8")
    if 'Extension="svg"' not in types:
        marker = '<Default Extension="rels"'
        assert marker in types, f"{path}: no rels default to insert beside"
        types = types.replace(
            marker,
            '<Default Extension="svg" ContentType="image/svg+xml"/>' + marker,
        )
    parts["[Content_Types].xml"] = types.encode("utf-8")


def rewrite(path: Path, parts: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def main() -> None:
    assert SVG.exists(), f"write {SVG} first (make_gif_svg_pictures.py)"
    write_fallback()
    write_docx(DOCX)
    write_pptx(PPTX)
    complete_with_svg(
        DOCX,
        "word/document.xml",
        "word/_rels/document.xml.rels",
        "word/media/image1.png",
        "word/media/image2.svg",
    )
    complete_with_svg(
        PPTX,
        "ppt/slides/slide1.xml",
        "ppt/slides/_rels/slide1.xml.rels",
        "ppt/media/image1.png",
        "ppt/media/image2.svg",
    )

    # The same document and drawing, with the drawing as the picture's own part.
    write_docx(EMBED_DOCX)
    write_pptx(EMBED_PPTX)
    embed_svg(
        EMBED_DOCX,
        "word/_rels/document.xml.rels",
        "word/media/image1.png",
        "word/media/diagram.svg",
    )
    embed_svg(
        EMBED_PPTX,
        "ppt/slides/_rels/slide1.xml.rels",
        "ppt/media/image1.png",
        "ppt/media/diagram.svg",
    )

    FALLBACK.unlink()
    for path in (DOCX, PPTX, EMBED_DOCX, EMBED_PPTX):
        print(f"wrote {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
