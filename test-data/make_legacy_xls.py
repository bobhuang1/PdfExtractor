#!/usr/bin/env python3
"""Write `test-data/ledger.xlsx` and `test-data/ledger.xls`.

The pair exists for the spreadsheet side of the LibreOffice bridge, the way
`legacy_deck.ppt` and `q3_review.pptx` exist for the presentation side: the
legacy file is the native one carried through a real legacy writer, so a test
can hold the whole bridged document to the native reading and fail on anything
the conversion changed. The native workbook is written here with `openpyxl`,
and the `.xls` is that workbook converted by LibreOffice's own "MS Excel 97"
filter — the same importer and exporter the bridge itself uses.

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_legacy_xls.py

Needs LibreOffice (`soffice` on `PATH`, or `PDFE_SOFFICE` naming it), the same
dependency the bridge itself has.

What the pair carries, one case per feature the spreadsheet reader has:

* `Ledger` — a header row of strings; numbers, a boolean and a merged range
  (`B5:C5`, a column span); an empty row between two blocks of cells, which
  leaves the sheet two tables rather than one; and `=SUM(B2:B3)`, whose cached
  result (`15`) is what a reader sees.
* `Notes` — a **row-spanning** merge (`A2:A3`, `row_span` two) beside the values
  it labels, and a date cell, which reads as `2024-03-05 00:00:00` rather than
  as the serial number the file stores.
* `Hidden` — a hidden worksheet, which the reader keeps on the `invisible`
  content layer instead of dropping. Calc's legacy writer keeps the sheet's
  hidden state, so the bridge has to carry the layer too.

The one thing the native file cannot say by itself is the formula's cached
result: `openpyxl` writes `<f>SUM(B2:B3)</f>` with an empty `<v>`, because it
does not calculate. Left that way, the cell is empty to a reader, and the
legacy side — where Calc *did* calculate and store `15` — would differ from the
native one for a reason that is the fixture's, not the bridge's. The cached
value is therefore written into the saved package, exactly as Excel writes it,
so both sides of the pair say `15`.
"""

import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
NATIVE = HERE / "ledger.xlsx"
BRIDGED = HERE / "ledger.xls"

# The formula cell, and the result a spreadsheet application caches for it.
CACHED_VALUES = {"B6": "15"}


def soffice() -> str:
    explicit = os.environ.get("PDFE_SOFFICE")
    if explicit:
        return explicit
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if not found:
        sys.exit("no soffice on PATH and PDFE_SOFFICE is not set")
    return found


def cache_formula_values(path: Path) -> None:
    """Put each formula's result into its cell, as Excel and Calc both do."""
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    sheet = parts["xl/worksheets/sheet1.xml"].decode("utf-8")
    for reference, value in CACHED_VALUES.items():
        pattern = re.compile(
            rf'(<c r="{reference}"[^>]*>\s*<f>[^<]*</f>)(?:<v>[^<]*</v>)?\s*(</c>)'
        )
        sheet, count = pattern.subn(rf"\g<1><v>{value}</v>\g<2>", sheet)
        assert count == 1, f"{reference}: cached value written {count} times"

    parts["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def write_native() -> None:
    workbook = Workbook()

    ledger = workbook.active
    ledger.title = "Ledger"
    ledger.append(["Item", "Qty", "Price", "Sold"])
    ledger.append(["Widget", 12, 1.5, True])
    ledger.append(["Gadget", 3, 22.75, False])
    ledger.append([])  # an empty row: the sheet's two data regions
    ledger["B5"] = "Shipping included"
    ledger.merge_cells("B5:C5")
    ledger["A6"] = "Total"
    ledger["B6"] = "=SUM(B2:B3)"

    notes = workbook.create_sheet("Notes")
    notes.append(["note", "value"])
    notes["A2"] = "owner"
    notes.merge_cells("A2:A3")
    notes["B2"] = "Ada"
    notes["B3"] = "Grace"
    notes.append(["updated", date(2024, 3, 5)])
    notes["B4"].number_format = "yyyy-mm-dd"

    hidden = workbook.create_sheet("Hidden")
    hidden.append(["internal"])
    hidden.append([42])
    hidden.sheet_state = "hidden"

    workbook.save(NATIVE)
    cache_formula_values(NATIVE)


def main() -> None:
    write_native()

    with tempfile.TemporaryDirectory(prefix="pdfe-legacy-xls-") as scratch:
        out = Path(scratch) / "out"
        out.mkdir()
        profile = Path(scratch) / "profile"
        result = subprocess.run(
            [
                soffice(),
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to",
                "xls",
                "--outdir",
                str(out),
                str(NATIVE),
            ],
            capture_output=True,
            text=True,
        )
        produced = out / f"{NATIVE.stem}.xls"
        if result.returncode != 0 or not produced.is_file():
            sys.exit(
                "LibreOffice could not write the fixture: "
                f"{result.stderr.strip() or result.stdout.strip() or 'no diagnostic'}"
            )
        shutil.copyfile(produced, BRIDGED)
    print(f"wrote {NATIVE} and {BRIDGED}")


if __name__ == "__main__":
    main()
