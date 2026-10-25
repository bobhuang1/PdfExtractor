#!/usr/bin/env python3
"""Author `master_background_deck.pptx`: a deck whose photograph is behind every
slide because the slide *master* states it, not the slides.

A background can be stated by the slide, by its layout, or by the layout's
master, and a part that states none inherits the one behind it. A deck's
photographic backdrop is usually the master's — one picture, every slide — so
the picture lives in `ppt/slideMasters/slideMaster1.xml`:

    <p:bg>
      <p:bgPr>
        <a:blipFill rotWithShape="1">
          <a:blip r:embed="rIdBackdrop"/>
          <a:stretch><a:fillRect/></a:stretch>
        </a:blipFill>
        <a:effectLst/>
      </p:bgPr>
    </p:bg>

naming the photograph through the master's own relationships. The master
python-pptx writes states a theme background instead
(`<p:bgRef idx="1001"><a:schemeClr val="bg1"/></p:bgRef>`), so the element is
replaced whole in the saved package; python-pptx cannot author a picture fill on
a background at any level.

The deck has three slides, all on the blank layout:

* the first two state no `p:bg` at all, so the photograph behind them is the
  master's — nothing in their own parts says so;
* the third states a colour background of its own, which is where the
  inheritance stops: a slide that says what is behind it is not given the
  master's photograph as well.

Each slide says a sentence, so the words in front of the backdrop are read
beside it.

Regenerate with:

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_master_background_deck.py

The deck is written by python-pptx — a real PPTX writer — and the master's
background is the one piece it cannot author, added the way PowerPoint writes
it. The tests pin what the reader makes of the package, never the container's
timestamps.
"""

from pathlib import Path
import zipfile

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
OUT = HERE / "master_background_deck.pptx"
BACKDROP = HERE / "images" / "scan.jpg"

MASTER = "ppt/slideMasters/slideMaster1.xml"
RELS = "ppt/slideMasters/_rels/slideMaster1.xml.rels"
MEDIA = "ppt/media/master_backdrop.jpeg"
REL_ID = "rIdBackdrop"
REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

BACKGROUND = (
    "<p:bg><p:bgPr>"
    f'<a:blipFill rotWithShape="1"><a:blip r:embed="{REL_ID}"/>'
    "<a:stretch><a:fillRect/></a:stretch></a:blipFill>"
    "<a:effectLst/>"
    "</p:bgPr></p:bg>"
)
COLOUR = (
    "<p:bg><p:bgPr>"
    '<a:solidFill><a:schemeClr val="accent2"/></a:solidFill>'
    "<a:effectLst/>"
    "</p:bgPr></p:bg>"
)


def main() -> None:
    deck = Presentation()

    for words in (
        "A photograph behind these words.",
        "The same photograph behind these words.",
        "A colour behind these words.",
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[6])  # blank
        box = slide.shapes.add_textbox(
            Inches(0.8), Inches(0.6), Inches(6.0), Inches(0.6)
        )
        box.text_frame.text = words

    deck.save(OUT)

    with zipfile.ZipFile(OUT) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    if MEDIA in parts:
        raise SystemExit(f"{MEDIA} is already in the deck")
    rels = parts[RELS].decode("utf-8")
    if REL_ID in rels:
        raise SystemExit(f"{RELS} already names {REL_ID}")
    parts[RELS] = rels.replace(
        "</Relationships>",
        f'<Relationship Id="{REL_ID}" Type="{REL_TYPE}" Target="../media/master_backdrop.jpeg"/>'
        "</Relationships>",
    ).encode("utf-8")
    parts[MEDIA] = BACKDROP.read_bytes()

    # The package already declares `jpeg` (python-pptx's default), but keep the
    # fixture honest if that changes: the part is a JPEG and must say so.
    types = parts["[Content_Types].xml"].decode("utf-8")
    if 'Extension="jpeg"' not in types:
        types = types.replace(
            "</Types>",
            '<Default Extension="jpeg" ContentType="image/jpeg"/></Types>',
        )
        parts["[Content_Types].xml"] = types.encode("utf-8")

    master = parts[MASTER].decode("utf-8")
    start = master.find("<p:bg>")
    end = master.find("</p:bg>")
    if start < 0 or end < 0:
        raise SystemExit(f"{MASTER} states no background to replace")
    parts[MASTER] = (
        master[:start] + BACKGROUND + master[end + len("</p:bg>") :]
    ).encode("utf-8")

    # The third slide says what is behind it, so the master's photograph must not
    # reach it: the colour is added the same way the master's picture is.
    third = "ppt/slides/slide3.xml"
    slide = parts[third].decode("utf-8")
    if "<p:bg>" in slide:
        raise SystemExit(f"{third} already has a background")
    marker = "<p:cSld>"
    if marker not in slide:
        raise SystemExit(f"{third} has no p:cSld")
    parts[third] = slide.replace(marker, marker + COLOUR, 1).encode("utf-8")

    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)

    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
