#!/usr/bin/env python3
"""Write `mc_svg_pictures.docx` and `mc_svg_pictures.pptx`.

A picture can be stated twice without an `asvg:svgBlip`: the modern drawing in
an `mc:AlternateContent`'s `mc:Choice`, and, in `mc:Fallback`, the copy an older
reader shows — for Word that is a `w:pict` with a VML `v:imagedata`, for
PowerPoint another `p:pic` whose blip is the raster. Only one of the two is the
picture of the document: a reader that takes the choice reports the drawing, one
that takes both reports the picture twice, and one that looks only at the
fallback reports a raster of it.

Each package holds two such pictures and a sentence before and after them:

* the first states the SVG in its choice and a grey box in its fallback — the
  drawing is `images/diagram.svg`, the 40 × 24 one the other SVG fixtures use,
  and the fallback is deliberately unlike it, so taking the wrong branch is
  visible in what comes out;
* the second's choice names a drawing the package does not hold, so the fallback
  — an orange 60 × 36 raster — is the picture an old reader would have shown and
  the one the document means. That is what a fallback is for.

Neither office library writes this form, and LibreOffice writes the other one
(the `asvg:svgBlip` extension, see `make_svg_blip_pictures.py`), so each package
is written with a raster in the picture's place and then completed the way a
converter completes it: the drawing part and the relationship naming it added
beside the fallback, a content-type declaration for the extension, and the
`mc:AlternateContent` pair written around the picture's own drawing — choice
first, fallback second, in the order the markup-compatibility schema states.

Regenerate with:

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_mc_svg_pictures.py

The libraries are used only to write the fixtures; the readers under test never
see this script, and they read the committed bytes.
"""

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt as PptxPt

HERE = Path(__file__).resolve().parent
IMAGES = HERE / "images"
SVG = IMAGES / "diagram.svg"
DOCX = HERE / "mc_svg_pictures.docx"
PPTX = HERE / "mc_svg_pictures.pptx"

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
V = "urn:schemas-microsoft-com:vml"
O = "urn:schemas-microsoft-com:office:office"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"

# What the two pictures are, in points: the drawing's own grid is 40 × 24, and
# the second picture is larger so the two rasters cannot be confused.
FIRST = (40.0, 24.0)
SECOND = (60.0, 36.0)

BOX = HERE / ".mc-svg-box.png"
ORANGE = HERE / ".mc-svg-orange.png"


def write_fallbacks() -> None:
    """The two rasters the packages carry as their fallbacks."""
    box = Image.new("RGB", (int(FIRST[0]), int(FIRST[1])), (128, 128, 128))
    ImageDraw.Draw(box).rectangle(
        [(2, 2), (int(FIRST[0]) - 3, int(FIRST[1]) - 3)], outline=(64, 64, 64)
    )
    box.save(BOX)

    orange = Image.new("RGB", (int(SECOND[0]), int(SECOND[1])), (224, 140, 40))
    ImageDraw.Draw(orange).rectangle(
        [(4, 4), (int(SECOND[0]) - 5, int(SECOND[1]) - 5)], outline=(120, 70, 20)
    )
    orange.save(ORANGE)


def write_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Before the pictures.")
    document.add_picture(str(BOX), width=Pt(FIRST[0]), height=Pt(FIRST[1]))
    document.add_picture(str(ORANGE), width=Pt(SECOND[0]), height=Pt(SECOND[1]))
    document.add_paragraph("After the pictures.")
    document.save(path)


def write_pptx(path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])  # blank

    def words(top: float) -> None:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(4.0), Inches(0.4))
        box.text_frame.text = f"{'Before' if top < 1 else 'After'} the pictures."

    words(0.4)
    slide.shapes.add_picture(
        str(BOX), Inches(0.8), Inches(1.0), width=PptxPt(FIRST[0]), height=PptxPt(FIRST[1])
    )
    slide.shapes.add_picture(
        str(ORANGE), Inches(0.8), Inches(2.0), width=PptxPt(SECOND[0]), height=PptxPt(SECOND[1])
    )
    words(3.5)
    deck.save(path)


def relationship_target(from_part: str, to_part: str) -> str:
    """`to_part`, as a relationship inside `from_part` names it."""
    assert to_part.startswith(("word/", "ppt/")), to_part
    if from_part.startswith("word/"):
        return to_part[len("word/") :]
    # A slide part's relationships resolve against `ppt/slides/`.
    return "../" + to_part[len("ppt/") :]


def add_part(
    parts: dict[str, bytes],
    drawing_part: str,
    rels_part: str,
    media_part: str,
    data: bytes,
    extension: str,
    mime: str,
) -> str:
    """Add a media part and the relationship naming it; return the relationship
    id, whether or not the part is written."""
    rels = parts[rels_part].decode("utf-8")
    used = [int(match.group(1)) for match in re.finditer(r'Id="rId(\d+)"', rels)]
    rel = f"rId{max(used) + 1}"
    parts[rels_part] = rels.replace(
        "</Relationships>",
        f'<Relationship Id="{rel}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
        f'Target="{relationship_target(drawing_part, media_part)}"/></Relationships>',
    ).encode("utf-8")
    if data:
        parts[media_part] = data

    types = parts["[Content_Types].xml"].decode("utf-8")
    if f'Extension="{extension}"' not in types:
        marker = '<Default Extension="rels"'
        assert marker in types, f"no rels default to insert beside the {extension} one"
        types = types.replace(
            marker, f'<Default Extension="{extension}" ContentType="{mime}"/>' + marker
        )
    parts["[Content_Types].xml"] = types.encode("utf-8")
    return rel


def declare(xml: str, root: str, prefixes: dict[str, str]) -> str:
    """Add any of `prefixes` the root element `root` does not declare."""
    start = xml.index(root)
    end = xml.index(">", start)
    tag = xml[start:end]
    missing = "".join(
        f' xmlns:{prefix}="{uri}"'
        for prefix, uri in prefixes.items()
        if f"xmlns:{prefix}=" not in tag
    )
    return xml[:end] + missing + xml[end:]


def complete_docx(path: Path) -> None:
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    svg_rel = add_part(
        parts,
        "word/document.xml",
        "word/_rels/document.xml.rels",
        "word/media/diagram.svg",
        SVG.read_bytes(),
        "svg",
        "image/svg+xml",
    )
    # The second picture's choice names a drawing that is not in the package.
    gone_rel = add_part(
        parts,
        "word/document.xml",
        "word/_rels/document.xml.rels",
        "word/media/gone.svg",
        b"",
        "svg",
        "image/svg+xml",
    )

    document = parts["word/document.xml"].decode("utf-8")
    document = declare(document, "<w:document", {"mc": MC, "v": V, "o": O, "wps": WPS})

    drawings = re.findall(r"<w:drawing>.*?</w:drawing>", document, re.DOTALL)
    assert len(drawings) == 2, f"{path}: expected two drawings, found {len(drawings)}"

    for drawing, choice_rel, (width, height) in zip(drawings, [svg_rel, gone_rel], [FIRST, SECOND]):
        embed = re.search(r'r:embed="([^"]+)"', drawing)
        assert embed, f"{path}: a drawing with no blip"
        fallback_rel = embed.group(1)
        choice = drawing.replace(f'r:embed="{fallback_rel}"', f'r:embed="{choice_rel}"')
        document = document.replace(
            drawing,
            "<mc:AlternateContent>"
            '<mc:Choice Requires="wps">'
            + choice
            + "</mc:Choice>"
            + "<mc:Fallback><w:pict>"
            '<v:shape id="Picture" o:spid="_x0000_i1025" filled="f" stroked="f" '
            f'style="width:{width:g}pt;height:{height:g}pt">'
            f'<v:imagedata r:id="{fallback_rel}" o:title=""/>'
            "</v:shape></w:pict></mc:Fallback>"
            "</mc:AlternateContent>",
        )

    parts["word/document.xml"] = document.encode("utf-8")
    rewrite(path, parts)


def complete_pptx(path: Path) -> None:
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    slide_part = "ppt/slides/slide1.xml"
    rels_part = "ppt/slides/_rels/slide1.xml.rels"
    svg_rel = add_part(
        parts,
        slide_part,
        rels_part,
        "ppt/media/diagram.svg",
        SVG.read_bytes(),
        "svg",
        "image/svg+xml",
    )
    gone_rel = add_part(parts, slide_part, rels_part, "ppt/media/gone.svg", b"", "svg", "image/svg+xml")

    slide = parts[slide_part].decode("utf-8")
    slide = declare(slide, "<p:sld", {"mc": MC, "p14": P14})

    pictures = re.findall(r"<p:pic>.*?</p:pic>", slide, re.DOTALL)
    assert len(pictures) == 2, f"{path}: expected two pictures, found {len(pictures)}"

    for picture, choice_rel in zip(pictures, [svg_rel, gone_rel]):
        embed = re.search(r'r:embed="([^"]+)"', picture)
        assert embed, f"{path}: a picture with no blip"
        fallback_rel = embed.group(1)
        choice = picture.replace(f'r:embed="{fallback_rel}"', f'r:embed="{choice_rel}"')
        slide = slide.replace(
            picture,
            "<mc:AlternateContent>"
            '<mc:Choice Requires="p14">'
            + choice
            + "</mc:Choice><mc:Fallback>"
            + picture
            + "</mc:Fallback></mc:AlternateContent>",
        )

    parts[slide_part] = slide.encode("utf-8")
    rewrite(path, parts)


def rewrite(path: Path, parts: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)


def main() -> None:
    assert SVG.exists(), f"write {SVG} first (make_gif_svg_pictures.py)"
    write_fallbacks()
    write_docx(DOCX)
    write_pptx(PPTX)
    complete_docx(DOCX)
    complete_pptx(PPTX)
    BOX.unlink()
    ORANGE.unlink()
    for path in (DOCX, PPTX):
        print(f"wrote {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
