#!/usr/bin/env python3
"""Write `svg_stretch_picture.docx` and its drawing, `images/stretch.svg`.

An SVG is authored against its own viewport and placed in whatever frame the
document gives it. When the two have different proportions, the drawing's root
`preserveAspectRatio` decides how they meet: the default `xMidYMid meet` keeps
the drawing's shape and leaves the rest of the frame transparent, `slice` fills
the frame and is clipped by it, and `none` stretches the two axes
independently. The other SVG fixtures here state no such attribute, so they all
ask the default; this one asks `none`.

The drawing is a 40 x 24 box split into two halves, red on the left and blue on
the right, and the package places it in a 2.5 x 0.5 in frame — five times as
wide as it is tall, where the drawing is five to three. A reader that ignored
the attribute (or always used the default) would centre the drawing at its own
shape: 120 x 72 pixels of ink with 120 pixels of margin on either side. Stated
as `none`, the drawing fills the frame corner to corner: 360 x 72 pixels, red on
the left half and blue on the right. Both readings are visible in the pixels
alone, which is what the test holds.

python-docx cannot place an SVG, so the picture is written with a one-pixel PNG
stub and completed the way LibreOffice stores a drawing — the media part
renamed to `image1.svg`, its bytes replaced, the picture's relationship pointed
at the new name, and the extension declared. The part name is asserted rather
than guessed, so a library that names it differently fails the generator instead
of leaving a PNG where the test says SVG.

Regenerate with:

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_svg_stretch_docx.py

The libraries are used only to write the fixture; the reader under test never
sees this script, and it reads the committed bytes.
"""

import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

HERE = Path(__file__).resolve().parent
IMAGES = HERE / "images"
SVG = IMAGES / "stretch.svg"
DOCX = HERE / "svg_stretch_picture.docx"

# The drawing: two halves, so a stretch that keeps the halves left and right is
# distinguishable from a crop or a scale that moved them.
SVG_BYTES = (
    '<svg xmlns="http://www.w3.org/2000/svg" width="40" height="24" '
    'viewBox="0 0 40 24" preserveAspectRatio="none">'
    '<rect width="20" height="24" fill="#c0392b"/>'
    '<rect x="20" width="20" height="24" fill="#2f6db5"/>'
    "</svg>"
).encode("utf-8")

# The frame the package places the drawing in: five times as wide as it is tall.
FRAME = (Inches(2.5), Inches(0.5))

# A one-pixel PNG, written only so the office library has something to place.
STUB = HERE / ".stretch-stub.png"


def write_drawing() -> None:
    IMAGES.mkdir(exist_ok=True)
    SVG.write_bytes(SVG_BYTES)
    Image.new("RGB", (1, 1), "white").save(STUB)


def write_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Before the picture.")
    document.add_picture(str(STUB), width=FRAME[0], height=FRAME[1])
    document.add_paragraph("After the picture.")
    document.save(path)


def replace_stub_with_svg(path: Path) -> None:
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    stub = "image1.png"
    svg = "image1.svg"
    media = [name for name in parts if name.endswith("/" + stub)]
    assert len(media) == 1, f"{path}: expected one {stub}, found {media}"
    media = media[0]
    svg_part = media[: -len(stub)] + svg

    parts[svg_part] = SVG_BYTES
    del parts[media]

    # Only relationships name a media part; the drawing refers to the
    # relationship, not to the file.
    for name in [name for name in parts if name.endswith(".rels")]:
        text = parts[name].decode("utf-8")
        parts[name] = text.replace(stub, svg).encode("utf-8")

    types = parts["[Content_Types].xml"].decode("utf-8")
    if 'Extension="svg"' not in types:
        anchor = '<Default Extension="rels"'
        assert anchor in types, f"{path}: no rels default to insert beside"
        types = types.replace(
            anchor,
            '<Default Extension="svg" ContentType="image/svg+xml"/>' + anchor,
        )
    parts["[Content_Types].xml"] = types.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)

    # The picture's own relationship is the one that names the drawing.
    rels = parts["word/_rels/document.xml.rels"].decode("utf-8")
    assert "image1.svg" in rels, f"{path}: no relationship names the drawing"


def main() -> None:
    write_drawing()
    write_docx(DOCX)
    replace_stub_with_svg(DOCX)
    STUB.unlink()
    for path in (SVG, DOCX):
        print(f"wrote {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
