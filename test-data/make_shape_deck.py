#!/usr/bin/env python3
"""Author `shape_deck.pptx`: a slide whose shapes PowerPoint paints itself.

A PowerPoint shape is not a picture in the package. An autoshape, a connector or
a group of them is markup — a `p:spPr` with a preset outline, a transform and a
fill and line that usually come from the theme — and PowerPoint draws it when
the slide is shown. A reader that only looked for an image relationship behind a
`p:pic` left every such shape out of the document, where a slide plainly shows a
box, an oval, a star and a line.

The deck is built with python-pptx and holds, on one slide:

* a rounded box, `120 × 60` pt, filled `#156082` with a 2 pt `#0F4761` outline,
  and the words *A box with words.* inside it: the shape is a picture and its
  words are the slide's text, both;
* a group in its own coordinate space — a rectangle outline and an oval filled
  `#6DA33A` — so a child's transform has to be mapped through the group's frame,
  and the group is one picture rather than one per shape;
* a five-pointed star with no colour of its own: python-pptx writes it a
  `p:style` whose `a:fillRef idx="1"` names `accent1`, so it only paints at all
  if the theme in `ppt/theme/theme1.xml` is read;
* a straight connector from the box to the star, drawn with a 2 pt blue line;
* a text box with no fill and no line — words, and no ink, so no picture;
* an embedded PNG, to hold the other half: a shape that *does* have an image
  part behind it is still the picture of that part, not a painted shape.

Regenerate with:

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_shape_deck.py

python-pptx writes the package — the layout, the master and the theme included —
so the fixture is a deck a real writer produced rather than parts composed here.
The ZIP timestamps are not fixed, so a regenerated package is not byte-identical;
the tests pin what the reader makes of it, never the container.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "shape_deck.pptx"

BLANK = 6  # the blank layout: no placeholders to inherit anything from


def main() -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[BLANK])

    words = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.4))
    words.text_frame.text = "Before the shapes."

    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1.5), Pt(120), Pt(60)
    )
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x15, 0x60, 0x82)
    box.line.color.rgb = RGBColor(0x0F, 0x47, 0x61)
    box.line.width = Pt(2)
    box.text_frame.text = "A box with words."

    # A group has a coordinate space of its own: python-pptx gives it a child
    # frame and maps what its children state through the frame it occupies.
    group = slide.shapes.add_group_shape()
    outline = group.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Emu(100 * 9144), Emu(50 * 9144)
    )
    outline.fill.background()
    outline.line.color.rgb = RGBColor(0x2F, 0x6D, 0xB5)
    oval = group.shapes.add_shape(
        MSO_SHAPE.OVAL, Emu(60 * 9144), Emu(0), Emu(40 * 9144), Emu(50 * 9144)
    )
    oval.fill.solid()
    oval.fill.fore_color.rgb = RGBColor(0x6D, 0xA3, 0x3A)
    oval.line.fill.background()
    group.left, group.top = Inches(1), Inches(3)
    group.width, group.height = Emu(100 * 9144), Emu(50 * 9144)

    # No fill is set: the colour is the `p:style` python-pptx writes, which
    # names `accent1` in the theme.
    star = slide.shapes.add_shape(
        MSO_SHAPE.STAR_5_POINT, Inches(5.5), Inches(1.5), Pt(80), Pt(80)
    )
    star.text_frame.text = ""

    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(2.5), Inches(2.2), Inches(5.5), Inches(2.2)
    )
    line.line.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    line.line.width = Pt(2)

    no_ink = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(5), Inches(0.4))
    no_ink.text_frame.text = "Words with no ink behind them."

    slide.shapes.add_picture(
        str(HERE / "invoice_scan.png"), Inches(6.5), Inches(3.5), width=Inches(3)
    )

    deck.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
