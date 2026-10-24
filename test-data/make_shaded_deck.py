#!/usr/bin/env python3
"""Author `shaded_deck.pptx`: a slide whose shapes are shaded, not flat.

A DrawingML shape is more than an outline and a solid colour. Its fill can be a
gradient measured across the shape, its box can cast a soft shadow, and a
connector's line can carry an arrowhead at either end. Each is markup in the
`p:spPr` — `a:gradFill`, `a:effectLst/a:outerShdw`, `a:ln/a:headEnd` — and
PowerPoint paints them when the slide is shown, so a reader that stops at a
solid fill and a plain line draws a slide that is not the one on screen.

The deck is built with python-pptx and holds, on one slide:

* a box, `180 × 90` pt, filled with a horizontal `#FF0000 → #0000FF` gradient
  and no line: the ramp has to be read, not its first stop;
* a box, `120 × 60` pt, filled `#6DA33A`, casting a soft `#000000` shadow at
  40 % — the shadow is drawn under the shape and reaches past it, so both the
  canvas and the ink are more than the box;
* two straight connectors, `120` pt long and 2 pt wide, one with a triangle at
  its tail, one with a stealth head at each end: the heads are ink at the line's
  ends, and the line stops at their base;
* a text box with no fill and no line — words, and no ink, so no picture.

Regenerate with:

    ~/.venvs/pdfextractor-docling/bin/python test-data/make_shaded_deck.py

python-pptx writes the package — the layout, the master and the theme included —
and states the gradient itself; the shadow and the arrowheads are the markup it
has no API for, so they are written into the shape's properties directly, with
the same attribute names and units the applications write. The ZIP timestamps
are not fixed, so a regenerated package is not byte-identical; the tests pin
what the reader makes of it, never the container.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "shaded_deck.pptx"

BLANK = 6  # the blank layout: no placeholders to inherit anything from

SHADOW = (
    f'<a:outerShdw {nsdecls("a")} blurRad="76200" dist="63500" '
    'dir="2700000" rotWithShape="0"><a:srgbClr val="000000">'
    '<a:alpha val="40000"/></a:srgbClr></a:outerShdw>'
)


def arrowhead(line, name: str, kind: str) -> None:
    """Put an arrowhead of `kind` on one end of a connector's line."""
    ln = line._get_or_add_ln()
    ln.append(
        parse_xml(
            f'<a:{name} {nsdecls("a")} type="{kind}" w="med" len="med"/>'
        )
    )


def shadowed(shape) -> None:
    """Give a shape the soft shadow PowerPoint draws for a raised box."""
    # `inherit = False` is python-pptx's way of writing the `a:effectLst`; the
    # effect itself goes inside that element, not beside it.
    shape.shadow.inherit = False
    effects = shape._element.spPr.find(qn("a:effectLst"))
    effects.append(parse_xml(SHADOW))


def main() -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[BLANK])

    # A two-stop ramp, left to right: python-pptx states the stops and the
    # `a:lin` that runs them across the shape.
    ramp = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Pt(60), Pt(60), Pt(180), Pt(90)
    )
    ramp.fill.gradient()
    ramp.fill.gradient_angle = 0.0
    first, last = ramp.fill.gradient_stops[0], ramp.fill.gradient_stops[-1]
    first.position, first.color.rgb = 0.0, RGBColor(0xFF, 0x00, 0x00)
    last.position, last.color.rgb = 1.0, RGBColor(0x00, 0x00, 0xFF)
    ramp.line.fill.background()

    shade = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Pt(300), Pt(60), Pt(120), Pt(60)
    )
    shade.fill.solid()
    shade.fill.fore_color.rgb = RGBColor(0x6D, 0xA3, 0x3A)
    shade.line.fill.background()
    shadowed(shade)

    tail = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Pt(60), Pt(240), Pt(180), Pt(240)
    )
    tail.line.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    tail.line.width = Pt(2)
    arrowhead(tail.line, "tailEnd", "triangle")

    both = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Pt(240), Pt(240), Pt(360), Pt(240)
    )
    both.line.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    both.line.width = Pt(2)
    arrowhead(both.line, "headEnd", "stealth")
    arrowhead(both.line, "tailEnd", "stealth")

    words = slide.shapes.add_textbox(Pt(60), Pt(360), Pt(300), Pt(20))
    words.text_frame.text = "Gradients, shadows and arrowheads."

    deck.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
